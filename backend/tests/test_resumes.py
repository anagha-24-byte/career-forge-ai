import io
import docx
import pytest


def create_sample_docx_bytes(text_content: str = None) -> bytes:
    doc = docx.Document()
    doc.add_heading("Anagha Sharma", 0)
    doc.add_paragraph("Email: student@college.edu | Phone: +91 9876543210")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.Tech in Computer Science and Engineering, 2023 - 2027. CGPA: 8.8/10")
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Languages: Python, JavaScript, SQL, C++. Frameworks: FastAPI, React, Tailwind CSS. Tools: Git, Docker.")
    doc.add_heading("Projects", level=1)
    content = text_content or (
        "CareerForge AI: Built a full-stack student placement platform using FastAPI, PostgreSQL, and React. "
        "Integrated AI resume analysis and mock interviews with Claude API. Benchmarked 1,000 requests."
    )
    doc.add_paragraph(content)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def test_upload_resume_docx_success(auth_client_a):
    file_bytes = create_sample_docx_bytes()
    res = auth_client_a.post(
        "/api/resumes",
        files={"file": ("resume.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["filename"] == "resume.docx"
    assert data["char_count"] > 100
    assert "FastAPI" in data["extracted_text"]


def test_upload_resume_invalid_format_rejected(auth_client_a):
    res = auth_client_a.post(
        "/api/resumes",
        files={"file": ("malicious.exe", b"binarycontent", "application/octet-stream")},
    )
    assert res.status_code == 400
    assert "Unsupported document format" in res.json()["detail"]


def test_list_resumes(auth_client_a):
    file_bytes = create_sample_docx_bytes()
    auth_client_a.post(
        "/api/resumes",
        files={"file": ("resume_alpha.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    res = auth_client_a.get("/api/resumes")
    assert res.status_code == 200
    assert len(res.json()) >= 1


def test_analyze_resume_success(auth_client_a):
    file_bytes = create_sample_docx_bytes()
    upload_res = auth_client_a.post(
        "/api/resumes",
        files={"file": ("my_resume.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    resume_id = upload_res.json()["id"]

    analyze_res = auth_client_a.post(
        f"/api/resumes/{resume_id}/analyze",
        json={"target_role": "Backend Engineer"},
    )
    assert analyze_res.status_code == 200
    data = analyze_res.json()
    assert "overall_score" in data
    assert 0 <= data["overall_score"] <= 100
    analysis = data["analysis_json"]
    assert "rubric_breakdown" in analysis
    assert len(analysis["rubric_breakdown"]) >= 3
    assert "ats_readability" in analysis
    assert "weak_bullet_points" in analysis
    assert "priority_recommendations" in analysis


def test_get_resume_analyses(auth_client_a):
    file_bytes = create_sample_docx_bytes()
    upload_res = auth_client_a.post(
        "/api/resumes",
        files={"file": ("doc_v1.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    resume_id = upload_res.json()["id"]

    auth_client_a.post(
        f"/api/resumes/{resume_id}/analyze",
        json={"target_role": "Full Stack Intern"},
    )

    analyses_res = auth_client_a.get(f"/api/resumes/{resume_id}/analyses")
    assert analyses_res.status_code == 200
    assert len(analyses_res.json()) >= 1
