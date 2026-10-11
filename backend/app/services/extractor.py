import io
import re
from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader
from docx import Document
from app.core.config import settings
from app.core.security import sanitize_untrusted_text


async def extract_text_from_file(file: UploadFile) -> tuple[str, int]:
    """
    Validates uploaded resume document (PDF or DOCX), enforces size limits,
    extracts plain text safely, and checks for minimum readable content.
    Returns: (extracted_text, file_size_bytes)
    """
    content = await file.read()
    file_size = len(content)

    if file_size > settings.MAX_UPLOAD_SIZE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File exceeds maximum allowed size of {settings.MAX_UPLOAD_SIZE_BYTES // (1024 * 1024)}MB.",
        )

    filename = (file.filename or "").lower()
    extracted_text = ""

    if filename.endswith(".pdf"):
        try:
            reader = PdfReader(io.BytesIO(content))
            pages_text = []
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                pages_text.append(page_text)
            extracted_text = "\n".join(pages_text)
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Could not read PDF document: {str(exc)}",
            )
    elif filename.endswith(".docx"):
        try:
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            extracted_text = "\n".join(paragraphs)
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Could not read DOCX document: {str(exc)}",
            )
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported document format. Only .pdf and .docx files are accepted.",
        )

    # Clean and normalize extracted text
    cleaned_text = re.sub(r"\s+", " ", extracted_text).strip()

    if len(cleaned_text) < 40:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded resume is empty or could not be parsed into readable text.",
        )

    sanitized = sanitize_untrusted_text(cleaned_text)
    return sanitized, file_size
